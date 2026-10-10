"""Local knowledge base: linked folders and past meetings, embedded with
EmbeddingGemma 2 and searched by cosine similarity in memory."""

from __future__ import annotations

import logging
import os
import threading
from pathlib import Path
from typing import Iterator

import numpy as np

from . import config, models
from .db import DB
from .events import bus
from .llama import runtime
from .notes import chars_per_token, transcript_text

log = logging.getLogger(__name__)

EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx", ".csv", ".json", ".html", ".htm", ".rtf", ".log"}
MAX_FILE_BYTES = 25 * 1024 * 1024
CHUNK_CHARS = 1200
OVERLAP_CHARS = 200
BATCH = 16
QUERY_PREFIX = "task: search result | query: "
MEETING_CONTEXT_TOKENS = 23000  # a context slot holds 32k; leaves room for the question and answer


def doc_prefix(title: str) -> str:
    return f"title: {title or 'none'} | text: "


def extract(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext in {".html", ".htm"}:
        import html
        import re

        raw = path.read_text(encoding="utf-8", errors="replace")
        raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
        return html.unescape(re.sub(r"<[^>]+>", " ", raw))
    return path.read_text(encoding="utf-8", errors="replace")


def _pieces(text: str, size: int) -> list[str]:
    """Paragraphs, with any paragraph longer than `size` split at whitespace."""
    out = []
    for para in text.replace("\r\n", "\n").split("\n\n"):
        para = " ".join(para.split())
        while len(para) > size:
            cut = para.rfind(" ", 0, size)
            cut = cut if cut > size // 2 else size
            out.append(para[:cut].strip())
            para = para[cut:].strip()
        if para:
            out.append(para)
    return out


def chunk(text: str, size: int = CHUNK_CHARS, overlap: int = OVERLAP_CHARS) -> list[str]:
    """Greedy paragraph packing into ~`size`-char chunks; each chunk starts with the
    last ~`overlap` chars of the previous one so facts spanning a boundary stay findable."""
    out: list[str] = []
    cur = ""
    for p in _pieces(text, size - overlap):
        if cur and len(cur) + 2 + len(p) > size:
            out.append(cur)
            tail = cur[-overlap:]
            cur = tail[tail.find(" ") + 1:] if " " in tail else tail
            cur = f"{cur} … {p}" if cur else p
        else:
            cur = f"{cur}\n\n{p}" if cur else p
    if cur:
        out.append(cur)
    return out


def _embed(texts: list[str]) -> np.ndarray:
    vecs = []
    for i in range(0, len(texts), BATCH):
        vecs.extend(runtime.embed(texts[i:i + BATCH]))
    m = np.asarray(vecs, dtype=np.float32)
    m /= np.linalg.norm(m, axis=1, keepdims=True) + 1e-12
    return m


def available() -> bool:
    return models.is_installed("embeddinggemma-2")


class Index:
    def __init__(self, db: DB):
        self.db = db
        self._lock = threading.Lock()
        self._meta: list[dict] | None = None
        self._mat: np.ndarray | None = None
        self.indexing = False
        self._scan_lock = threading.Lock()

    def _invalidate(self) -> None:
        self._meta = self._mat = None

    def _load(self) -> tuple[list[dict], np.ndarray]:
        if self._meta is None:
            self._meta, self._mat = self.db.kb_all()
        return self._meta, self._mat

    def index_text(self, ref: str, title: str, text: str, file_meta: tuple[float, int] | None = None) -> int:
        pieces = chunk(text)
        if not pieces:
            self.db.kb_replace(ref, title, [], np.zeros((0, 1)), file_meta)
            return 0
        vecs = _embed([doc_prefix(title) + p for p in pieces])
        self.db.kb_replace(ref, title, pieces, vecs, file_meta)
        with self._lock:
            self._invalidate()
        return len(pieces)

    def index_meeting(self, mid: str) -> None:
        m = self.db.meeting(mid)
        if not m or not available():
            return
        body = "\n\n".join(x for x in (m["enhanced"] or m["notes"], transcript_text(self.db.lines(mid))) if x)
        self.index_text(f"meeting:{mid}", m["title"], body)

    def scan(self) -> None:
        """Index new or changed files in linked folders; drop files that disappeared."""
        if not available() or not self._scan_lock.acquire(blocking=False):
            return  # a scan is already running; it will pick up the latest folder list next time
        self.indexing = True
        try:
            folders = [Path(f) for f in config.settings().folders]
            seen: set[str] = set()
            todo: list[tuple[Path, os.stat_result]] = []
            for root in folders:
                if not root.is_dir():
                    continue
                for dirpath, dirnames, filenames in os.walk(root):
                    dirnames[:] = [d for d in dirnames if not d.startswith((".", "node_modules", "__pycache__"))]
                    for fn in filenames:
                        p = Path(dirpath) / fn
                        if p.suffix.lower() not in EXTENSIONS:
                            continue
                        try:
                            st = p.stat()
                        except OSError:
                            continue
                        if st.st_size > MAX_FILE_BYTES:
                            continue
                        key = str(p)
                        seen.add(key)
                        known = self.db.kb_file(key)
                        if known and known["mtime"] == st.st_mtime and known["size"] == st.st_size:
                            continue
                        todo.append((p, st))
            for f in self.db.kb_files():
                if f["path"] not in seen:
                    self.db.kb_remove(f["path"])
            for i, (p, st) in enumerate(todo):
                bus.publish("kb_progress", done=i, total=len(todo), current=p.name)
                try:
                    self.index_text(str(p), p.name, extract(p), (st.st_mtime, st.st_size))
                except Exception as e:
                    log.warning("could not index %s: %s", p, e)
                    # remember the failure so we don't retry an unreadable file every scan
                    self.db.kb_replace(str(p), p.name, [], np.zeros((0, 1)), (st.st_mtime, st.st_size))
            with self._lock:
                self._invalidate()
            bus.publish("kb_progress", done=len(todo), total=len(todo), current=None)
        finally:
            self.indexing = False
            self._scan_lock.release()

    def search(self, query: str, k: int = 8, refs: set[str] | None = None) -> list[dict]:
        with self._lock:
            meta, mat = self._load()
        if not meta:
            return []
        q = _embed([QUERY_PREFIX + query])[0]
        scores = mat @ q
        order = np.argsort(-scores)
        hits = []
        for i in order:
            if refs is not None and meta[i]["ref"] not in refs:
                continue
            hits.append({**meta[i], "score": float(scores[i])})
            if len(hits) == k:
                break
        return hits

    def stats(self) -> dict:
        files = self.db.kb_files()
        return {"files": len(files), "chunks": sum(f["chunks"] for f in files), "indexing": self.indexing}


ANSWER_SYSTEM = (
    "Answer the question using only the numbered sources. Cite sources inline like [1] or [2][3]. "
    "If the sources don't contain the answer, say so plainly and say what is missing. Be concise and specific."
)


def ask(index: Index, db: DB, question: str, meeting_id: str | None = None) -> tuple[list[dict], Iterator[str]]:
    """Returns (sources, token stream). With meeting_id, answers from that meeting's transcript and notes."""
    if meeting_id:
        m = db.meeting(meeting_id)
        if not m:
            raise KeyError(meeting_id)
        body = transcript_text(db.lines(meeting_id))
        notes = m["enhanced"] or m["notes"]
        text = (f"Notes:\n{notes}\n\n" if notes else "") + f"Transcript:\n{body}"
        ref = f"meeting:{meeting_id}"
        limit = int(MEETING_CONTEXT_TOKENS * chars_per_token(m["language"]))
        if len(text) > limit and available():
            # too long for one context slot: retrieve the relevant parts of this meeting only
            if m["status"] == "done" and not db.kb_has(ref):
                index.index_meeting(meeting_id)
            sources = index.search(question, k=12, refs={ref})
        else:
            sources = [{"ref": ref, "title": m["title"], "text": text[-limit:]}]
    else:
        if not available():
            raise RuntimeError("EmbeddingGemma 2 is not downloaded, so linked folders can't be searched")
        sources = index.search(question)
    context = "\n\n".join(f"[{i + 1}] {s['title']}\n{s['text']}" for i, s in enumerate(sources))
    msgs = [{"role": "system", "content": ANSWER_SYSTEM},
            {"role": "user", "content": f"Sources:\n{context or '(no sources found)'}\n\nQuestion: {question}"}]
    public = [{"n": i + 1, "ref": s["ref"], "title": s["title"], "excerpt": s["text"][:400],
               "score": s.get("score")} for i, s in enumerate(sources)]
    return public, runtime.chat_stream(msgs, max_tokens=1500, temperature=0.2)
